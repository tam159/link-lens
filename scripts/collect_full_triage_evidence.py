"""Collect full-shortlist audit evidence; does not assign relevance verdicts."""

import csv
import io
import itertools
import json
import subprocess
import tempfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader
from link_lens import store

ROOT = Path("outputs/ai-audit")
ROOT.mkdir(parents=True, exist_ok=True)
# Fresh evidence for the new shortlist; previous audit labels are never model inputs.
old = {}
shortlist = [
    json.loads(l) for l in Path("outputs/shortlist.jsonl").read_text().splitlines()
]


def collect(item):
    key = item["dataset_id"]
    if key in old:
        return {
            **old[key],
            "evidence_reused_from": "2026-09-18 audit; re-assessed 2026-09-19",
        }
    d = store.require("sources", key)["metadata"]
    result = {
        "item_id": key,
        "title": d["title"],
        "notes": d.get("notes", ""),
        "resources": d.get("resources", []),
        "attempts": [],
        "collected_at": store.now(),
    }
    choices = [
        r
        for r in result["resources"]
        if any(t in r.get("format", "").lower() for t in ["csv", "xlsx", "json", "pdf"])
    ]

    # Notices identify the notifying company in their PDF, not always the site spreadsheet.
    def priority(r):
        f = r.get("format", "").lower()
        if "anticipatory" in d["title"].lower() and "pdf" in f:
            return -1
        return 0 if "csv" in f else 1 if "xlsx" in f else 2 if "json" in f else 3

    choices.sort(key=priority)
    for res in choices[:2]:
        fmt = res.get("format", "").lower()
        a = {"url": res["url"], "resource_id": res["id"], "format": fmt}
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sample"
            cmd = [
                "curl",
                "--fail",
                "--location",
                "--silent",
                "--show-error",
                "--connect-timeout",
                "10",
                "--max-time",
                "40",
                "--max-filesize",
                "30000000",
                "-o",
                str(path),
            ]
            if "csv" in fmt:
                cmd += ["--range", "0-524287"]
            p = subprocess.run(cmd + [res["url"]], capture_output=True)
            a["exit_code"] = p.returncode
            if p.returncode:
                a["error"] = p.stderr.decode()[:300]
            else:
                data = path.read_bytes()
                a.update(
                    artifact_id=store.blob(
                        data, "application/octet-stream", "full-triage-sample"
                    ),
                    bytes=len(data),
                )
                try:
                    if not data:
                        raise ValueError("Empty resource")
                    if data.lstrip().lower().startswith((b"<!doctype html", b"<html")):
                        raise ValueError("HTML returned instead of data")
                    if "pdf" in fmt:
                        book = PdfReader(io.BytesIO(data))
                        a["pages"] = [
                            {"page": n + 1, "text": page.extract_text()}
                            for n, page in enumerate(book.pages[:3])
                        ]
                    elif "xlsx" in fmt:
                        book = load_workbook(
                            io.BytesIO(data), read_only=True, data_only=True
                        )
                        a["sheets"] = [
                            {
                                "sheet": s.title,
                                "rows": [
                                    [str(v) if v is not None else "" for v in row]
                                    for row in itertools.islice(
                                        s.iter_rows(values_only=True), 10
                                    )
                                ],
                            }
                            for s in book.worksheets
                        ]
                        book.close()
                    elif "json" in fmt:
                        a["json_sample"] = str(json.loads(data))[:18000]
                    else:
                        try:
                            text = data.decode("utf-8-sig")
                        except UnicodeDecodeError:
                            text = data.decode("cp1252")
                        dialect = csv.Sniffer().sniff(text[:16000], delimiters=",\t;|")
                        a["rows"] = list(
                            itertools.islice(
                                csv.reader(io.StringIO(text), dialect=dialect), 10
                            )
                        )
                except Exception as exc:
                    a["parse_error"] = str(exc)
        result["attempts"].append(a)
        if not a["exit_code"] and not a.get("parse_error"):
            break
    return result


with ThreadPoolExecutor(max_workers=4) as pool:
    rows = list(pool.map(collect, shortlist))
for n, row in enumerate(rows, 1):
    row["rank"] = n
(ROOT / "full-triage-evidence.json").write_text(json.dumps(rows, indent=2) + "\n")
for r in rows:
    print(
        r["rank"],
        r["title"],
        [(a["exit_code"], a.get("parse_error"), a.get("bytes")) for a in r["attempts"]],
    )
