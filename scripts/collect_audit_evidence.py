"""Collect evidence only; this script makes no evaluation verdicts."""

import csv, io, json, subprocess, tempfile
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
from openpyxl import load_workbook
from link_lens import store
from link_lens.readers import read_records
from link_lens.contracts import MappingSpec

root = Path("outputs/ai-audit")
root.mkdir(exist_ok=True)
packet = json.loads(Path("outputs/link-review.json").read_text())
batch = store.require("batches", packet["batch_id"])
result = store.read_json(batch["result_artifact"])
raw = {}
sources = {}
for rid in batch["run_ids"]:
    r = store.require("runs", rid)
    snap = store.require("snapshots", r["snapshot_id"])
    spec = MappingSpec.model_validate(
        store.require("mappings", r["mapping_id"])["config"]
    )
    _, rows = read_records(
        store.blob_path(snap["artifact_id"]),
        spec.reader,
        snap["sha256"],
        snap["receipt"].get("partial", False),
        25000,
    )
    sources[r["source_id"]] = {
        "slug": r["source_slug"],
        "snapshot": snap["sha256"],
        "url": snap["receipt"]["url"],
        "role": spec.subject_role,
        "filters": [f.model_dump() for f in spec.filters],
    }
    for row in rows:
        raw[(r["source_id"], row["record_id"])] = row
links = []
for n, item in enumerate(packet["items"], 1):
    e = item["evidence_to_review"]
    records = []
    for side in ["source_a_record", "source_b_record"]:
        ref = tuple(e[side])
        records.append({"source": sources[ref[0]], **raw[ref]})
    links.append(
        {"number": n, "item_id": item["item_id"], "link": e, "raw_records": records}
    )
(root / "link-evidence.json").write_text(json.dumps(links, indent=2))
triage = json.loads(Path("outputs/triage-review.json").read_text())


def collect(entry):
    d = store.require("sources", entry["item_id"])["metadata"]
    resources = d.get("resources", [])
    supported = [
        r
        for r in resources
        if r.get("format", "").lower() in ["csv", "tsv"]
        or "xlsx" in r.get("format", "").lower()
    ]
    supported.sort(
        key=lambda r: 0 if r.get("format", "").lower() in ["csv", "tsv"] else 1
    )
    result = {
        "item_id": entry["item_id"],
        "title": d["title"],
        "notes": d.get("notes", ""),
        "resources": resources,
        "attempts": [],
    }
    for r in supported[:2]:
        url = r["url"]
        fmt = r.get("format", "").lower()
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "data"
            cmd = [
                "curl",
                "--fail",
                "--location",
                "--silent",
                "--show-error",
                "--connect-timeout",
                "10",
                "--max-time",
                "40",
                "--max-filesize",
                "30000000",
                "-o",
                str(path),
            ]
            if fmt in ["csv", "tsv"]:
                cmd += ["--range", "0-524287"]
            p = subprocess.run(cmd + [url], capture_output=True)
            attempt = {"url": url, "format": fmt, "exit_code": p.returncode}
            if p.returncode:
                attempt["error"] = p.stderr.decode()[:300]
            else:
                data = path.read_bytes()
                sha = store.blob(data, "application/octet-stream", "audit-sample")
                attempt.update(artifact_id=sha, bytes=len(data), sample_only=True)
                try:
                    if "xlsx" in fmt:
                        book = load_workbook(
                            io.BytesIO(data), read_only=True, data_only=True
                        )
                        attempt["sheets"] = [
                            {
                                "sheet": s.title,
                                "rows": [
                                    [str(v) if v is not None else "" for v in row]
                                    for row in list(
                                        __import__("itertools").islice(
                                            s.iter_rows(values_only=True), 8
                                        )
                                    )
                                ],
                            }
                            for s in book.worksheets
                        ]
                        book.close()
                    else:
                        try:
                            text = data.decode("utf-8-sig")
                        except UnicodeDecodeError:
                            text = data.decode("cp1252", errors="replace")
                        try:
                            delimiter = (
                                csv.Sniffer()
                                .sniff(text[:16000], delimiters=",\t;|")
                                .delimiter
                            )
                        except csv.Error:
                            delimiter = ","
                        attempt["rows"] = list(
                            __import__("itertools").islice(
                                csv.reader(
                                    io.StringIO(text, newline=""), delimiter=delimiter
                                ),
                                8,
                            )
                        )
                except Exception as exc:
                    attempt["parse_error"] = str(exc)
            result["attempts"].append(attempt)
            if p.returncode == 0 and not attempt.get("parse_error"):
                break
    return result


with ThreadPoolExecutor(max_workers=4) as pool:
    results = list(pool.map(collect, triage["items"]))
(root / "triage-evidence.json").write_text(json.dumps(results, indent=2))
print(
    "Collected", len(links), "raw link pairs and", len(results), "dataset inspections"
)
