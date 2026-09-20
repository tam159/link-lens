"""Bounded GET + reader preflight for triage; never an entity-relevance label.

CSV checks establish sample availability, not successful full-file ingestion.
XLSX/JSON must fit the full-response bound. Each receipt is immutable evidence.
"""

import csv
import io
import ipaddress
import itertools
import socket
import time
from concurrent.futures import ThreadPoolExecutor
from urllib.parse import urlsplit, urlunsplit

import httpx
from openpyxl import load_workbook

from . import store

POLICY = {
    "version": "download-preflight-3",
    "csv_sample_bytes": 1_048_576,
    "full_resource_max_bytes": 100_000_000,
    "resource_seconds": 60,
    "max_resources_per_dataset": 12,
    "publisher_cap": 8,
    "scope": "At least one nonempty readable CSV/TSV sample or complete XLSX/JSON resource; not full-dataset validation or confirmed business relevance",
}


def parse_sample(data, fmt, partial=False):
    import json

    if not data.strip():
        raise ValueError("Empty response")
    if data.lstrip().lower().startswith((b"<!doctype html", b"<html")):
        raise ValueError("HTML returned instead of data")
    if fmt == "xlsx":
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            sheets = [
                {
                    "sheet": s.title,
                    "rows": [
                        [str(v) if v is not None else "" for v in row]
                        for row in itertools.islice(s.iter_rows(values_only=True), 20)
                    ],
                }
                for s in book.worksheets
            ]
            if not any(len(s["rows"]) >= 2 for s in sheets):
                raise ValueError("No worksheet with data rows")
            return {"sheets": sheets}
        finally:
            book.close()
    if fmt == "json":
        obj = json.loads(data)
        if not obj or not isinstance(obj, (list, dict)):
            raise ValueError("Empty or scalar JSON")
        return {"json_sample": str(obj)[:25000]}
    if data.startswith(b"PK"):
        raise ValueError("ZIP payload labelled CSV; archive reader not supported")
    try:
        text = data.decode("utf-8-sig")
        encoding = "utf-8-sig"
    except UnicodeDecodeError:
        text = data.decode("cp1252")
        encoding = "cp1252"
    if partial:
        # Drop possibly cut final physical line; quoted CSV remains checked below.
        text = text.rsplit("\n", 1)[0]
    lines = text.splitlines()
    header = next((line for line in lines if line.strip()), "")
    delimiter = max([",", "\t", ";", "|"], key=header.count)
    if not header.count(delimiter):
        raise ValueError("No tabular delimiter in header")
    rows = list(
        itertools.islice(
            csv.reader(io.StringIO(text, newline=""), delimiter=delimiter), 200
        )
    )
    if len(rows) < 2 or not any(len(r) > 1 and any(r) for r in rows[1:]):
        raise ValueError("No tabular data rows")
    if {c.strip().lower() for c in rows[0]} == {"resource name", "type", "download"}:
        raise ValueError(
            "Resource index points to other files; not a data table supported by this reader"
        )
    return {"encoding": encoding, "delimiter": delimiter, "rows": rows}


def public_resource_url(url):
    """Follow publisher-catalogued public hosting, including government cloud CDNs."""
    parts = urlsplit(url)
    if (
        parts.scheme not in ("http", "https")
        or parts.username
        or parts.password
        or parts.port not in (None, 80, 443)
    ):
        raise ValueError("Only public HTTP(S) resource URLs without credentials")
    if not parts.hostname:
        raise ValueError("Missing resource hostname")
    # Try the HTTPS equivalent of legacy HTTP catalogue links.
    url = urlunsplit(
        ("https", parts.netloc.removesuffix(":80"), parts.path, parts.query, "")
    )
    addresses = socket.getaddrinfo(parts.hostname, 443, type=socket.SOCK_STREAM)
    if not addresses or any(
        not ipaddress.ip_address(a[4][0]).is_global for a in addresses
    ):
        raise ValueError("Non-public destination blocked")
    return url


def probe_resource(resource):
    from .ingestion import normal_format

    fmt = normal_format(resource)
    receipt = {
        "resource_id": resource.get("id"),
        "url": resource.get("url"),
        "format": fmt,
        "checked_at": store.now(),
        "status": "unavailable",
    }
    started = time.monotonic()
    try:
        url = public_resource_url(resource["url"])
        limit = (
            POLICY["csv_sample_bytes"]
            if fmt == "csv"
            else POLICY["full_resource_max_bytes"]
        )
        headers = {"Range": f"bytes=0-{limit - 1}"} if fmt == "csv" else {}
        data = bytearray()
        partial = False
        # Redirects are checked before following; no private/local service access.
        with httpx.Client(timeout=20, follow_redirects=False) as client:
            for _ in range(6):
                url = public_resource_url(url)
                with client.stream("GET", url, headers=headers) as response:
                    if response.is_redirect:
                        url = str(response.url.join(response.headers["location"]))
                        continue
                    receipt.update(
                        http_status=response.status_code,
                        final_url=str(response.url),
                        content_type=response.headers.get("content-type"),
                    )
                    response.raise_for_status()
                    partial = response.status_code == 206
                    for chunk in response.iter_bytes():
                        if time.monotonic() - started > POLICY["resource_seconds"]:
                            raise TimeoutError("Resource elapsed-time budget")
                        remaining = limit - len(data)
                        data.extend(chunk[:remaining])
                        if len(data) >= limit:
                            if fmt != "csv":
                                raise ValueError(
                                    "Full resource exceeds preflight byte limit"
                                )
                            partial = True
                            break
                    break
            else:
                raise ValueError("Redirect limit exceeded")
        receipt.update(bytes=len(data), partial=partial)
        if data:
            receipt["artifact_id"] = store.blob(
                bytes(data), "application/octet-stream", "download-preflight-sample"
            )
        receipt["sample"] = parse_sample(bytes(data), fmt, partial)
        receipt["status"] = "available"
    except Exception as exc:  # noqa: BLE001 -- retain per-resource failures, continue triage
        receipt["error"] = f"{type(exc).__name__}: {str(exc)[:400]}"
    receipt["wall_seconds"] = round(time.monotonic() - started, 3)
    return receipt


def probe_dataset(dataset):
    from .ingestion import normal_format

    resources = [r for r in dataset.get("resources", []) if normal_format(r)]
    resources.sort(key=lambda r: {"csv": 0, "xlsx": 1, "json": 2}[normal_format(r)])
    result = {"dataset_id": dataset["id"], "attempts": [], "status": "unavailable"}
    for resource in resources[: POLICY["max_resources_per_dataset"]]:
        attempt = probe_resource(resource)
        result["attempts"].append(attempt)
        if attempt["status"] == "available":
            result.update(status="available", resource_id=attempt["resource_id"])
            break
    return result


def select_downloadable(
    ranked, datasets, *, check=probe_dataset, size=50, workers=4, receipts=None
):
    """Preserve rank order and publisher cap; promote no specific source IDs."""
    from .ingestion import normal_format

    receipts = {} if receipts is None else receipts
    selected, counts = [], {}
    eligible = [
        r
        for r in ranked
        if any(normal_format(x) for x in datasets[r["dataset_id"]].get("resources", []))
    ]
    for start in range(0, len(eligible), 12):
        group = [
            r
            for r in eligible[start : start + 12]
            if counts.get(r["publisher"], 0) < POLICY["publisher_cap"]
        ]
        todo = [r for r in group if r["dataset_id"] not in receipts]
        with ThreadPoolExecutor(max_workers=workers) as pool:
            for result in pool.map(check, [datasets[r["dataset_id"]] for r in todo]):
                receipts[result["dataset_id"]] = result
        for row in group:
            receipt = receipts[row["dataset_id"]]
            pub = row["publisher"]
            if (
                receipt["status"] != "available"
                or counts.get(pub, 0) >= POLICY["publisher_cap"]
            ):
                continue
            selected.append(
                {
                    **row,
                    "downloadability": {
                        "status": "sample_verified",
                        "resource_id": receipt["resource_id"],
                        "policy": POLICY["version"],
                    },
                    "reason": row["reason"]
                    + "; bounded GET and supported-reader preflight passed",
                }
            )
            counts[pub] = counts.get(pub, 0) + 1
            if len(selected) == size:
                return selected, receipts
    raise ValueError(
        f"Only {len(selected)} verified candidates under publisher cap; broaden queries rather than insert datasets manually"
    )
