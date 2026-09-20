"""Readers vary by format, never by publisher. Locators address immutable snapshots."""

import csv
import codecs
import hashlib
import io
import json
from datetime import date, datetime
from itertools import islice
from pathlib import Path
from openpyxl import load_workbook
from .contracts import ReaderSpec


def scalar(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def headers_for(values):
    headers = []
    for i, value in enumerate(values, 1):
        key = scalar(value).strip() or f"__unnamed_{i}"
        if key in headers:
            key = f"{key}__column_{i}"
        headers.append(key)
    return headers


def complete_prefix(data):
    # Prefixes may use CRLF, LF, or CR. csv.reader still checks quoted records.
    end = max(data.rfind(b"\n"), data.rfind(b"\r"))
    return data[: end + 1] if end >= 0 else b""


def read_records(
    path: Path, spec: ReaderSpec, snapshot_sha: str, partial=False, limit=5000
):
    if spec.format == "xlsx":
        book = load_workbook(
            io.BytesIO(path.read_bytes()), read_only=True, data_only=True
        )
        try:
            if spec.sheet not in book.sheetnames:
                raise ValueError(f"Select a real sheet: {book.sheetnames}")
            iterator = book[spec.sheet].iter_rows(values_only=True)
            for _ in range(spec.header_row - 1):
                next(iterator)
            headers = headers_for(next(iterator))
            records = []
            for n, values in enumerate(islice(iterator, limit), spec.header_row + 1):
                if not any(v is not None for v in values):
                    continue
                records.append(
                    record(
                        snapshot_sha,
                        spec.sheet,
                        n,
                        dict(zip(headers, map(scalar, values))),
                    )
                )
            return headers, records
        finally:
            book.close()
    if spec.format == "json":
        data = json.loads(path.read_bytes())
        for part in spec.json_records_path:
            data = data[part]
        if not isinstance(data, list):
            raise ValueError("JSON records path must resolve to an array")
        records = []
        headers = []
        for n, item in enumerate(data[:limit], 1):
            if not isinstance(item, dict):
                raise ValueError("JSON array must contain objects")
            row = {k: scalar(v) for k, v in flatten(item).items()}
            headers.extend(k for k in row if k not in headers)
            records.append(
                record(snapshot_sha, "json:" + ".".join(spec.json_records_path), n, row)
            )
        for row in records:
            row["values"] = {h: row["values"].get(h, "") for h in headers}
        return headers, records
    data = path.read_bytes()
    if partial:
        data = complete_prefix(data)
    text = data.decode(spec.encoding)
    iterator = csv.reader(
        io.StringIO(text, newline=""), delimiter=spec.delimiter, strict=True
    )
    for _ in range(spec.header_row - 1):
        next(iterator)
    headers = headers_for(next(iterator))
    records = []
    try:
        for n, values in enumerate(iterator, spec.header_row + 1):
            if not any(values):
                continue
            if len(values) != len(headers):
                raise ValueError(
                    f"Column count mismatch at logical row {n}: {len(values)} != {len(headers)}"
                )
            records.append(record(snapshot_sha, "csv", n, dict(zip(headers, values))))
            if len(records) >= limit:
                break
    except csv.Error as exc:
        if not partial or "unexpected end of data" not in str(exc):
            raise
        # Only the truncated final quoted record is discarded; preceding records remain.
    return headers, records


def flatten(item, prefix=""):
    result = {}
    for key, value in item.items():
        target = f"{prefix}.{key}" if prefix else key
        if isinstance(value, dict):
            result.update(flatten(value, target))
        else:
            result[target] = value
    return result


def record(sha, sheet, n, values):
    locator = {"snapshot_sha": sha, "sheet": sheet, "row": n}
    return {
        "record_id": hashlib.sha256(
            json.dumps(locator, sort_keys=True).encode()
        ).hexdigest(),
        "locator": locator,
        "values": values,
    }


def partition(rows):
    parts = {"discovery": [], "validation": [], "final": []}
    for row in rows:
        # All structural previews are confined to the first 10 rows, always discovery.
        bucket = int(row["record_id"][:8], 16) % 5
        name = (
            "discovery"
            if row["locator"]["row"] <= 10 or bucket < 3
            else ("validation" if bucket == 3 else "final")
        )
        parts[name].append(row)
    return parts


def inspect_resource(path: Path, format_hint: str, partial=False):
    """Only structural metadata and first 6 rows; never arbitrary body windows."""
    head = path.read_bytes()[:32]
    if head.startswith(b"PK"):
        book = load_workbook(
            io.BytesIO(path.read_bytes()), read_only=True, data_only=True
        )
        try:
            return {
                "format": "xlsx",
                "sheets": [
                    {
                        "name": s.title,
                        "preview": [
                            [scalar(v) for v in row]
                            for row in islice(s.iter_rows(values_only=True), 6)
                        ],
                    }
                    for s in book.worksheets
                ],
            }
        finally:
            book.close()
    if format_hint == "json" or head.lstrip().startswith((b"{", b"[")):
        data = json.loads(path.read_bytes())

        def shape(value, depth=0):
            if depth > 3:
                return type(value).__name__
            if isinstance(value, dict):
                return {k: shape(v, depth + 1) for k, v in value.items()}
            if isinstance(value, list):
                return {
                    "array_length": len(value),
                    "first_item": shape(value[0], depth + 1) if value else None,
                }
            return type(value).__name__

        return {"format": "json", "shape": shape(data)}
    encoding = "utf-8-sig"
    try:
        text = codecs.getincrementaldecoder(encoding)().decode(
            path.read_bytes(), final=not partial
        )[:32000]
    except UnicodeDecodeError:
        encoding = "cp1252"
        text = codecs.getincrementaldecoder(encoding)().decode(
            path.read_bytes(), final=not partial
        )[:32000]
    try:
        delimiter = csv.Sniffer().sniff(text[:8000], delimiters=",\t;|").delimiter
    except csv.Error:
        delimiter = ","
    try:
        preview = list(
            islice(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter), 6)
        )
    except csv.Error:
        preview = []
    return {
        "format": "csv",
        "detected_encoding": encoding,
        "detected_delimiter": delimiter,
        "preview": preview,
        "partial": partial,
    }
